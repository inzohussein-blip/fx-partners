#!/usr/bin/env python3
"""
Build the broker verification sheet (xlsx, Google-Sheets compatible) from a
JSON export of public.brokers.

    python3 scripts/broker-sheet/export_sheet.py brokers.json -o FX-Partners-Brokers-Data.xlsx

brokers.json is an array of broker rows with their broker_links, e.g. from:

    select json_agg(b order by b.sort_order, b.name) from (
      select br.*, coalesce((select json_agg(l) from public.broker_links l
                             where l.broker_id = br.id), '[]') as broker_links
      from public.brokers br) b;

The sheet keeps a hidden copy of every row («الأصل»); import_sheet.py diffs
against it, so only cells someone changed are written back.
"""
import argparse
import json
from datetime import date
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

_ap = argparse.ArgumentParser()
_ap.add_argument("json")
_ap.add_argument("-o", "--out", default="FX-Partners-Brokers-Data.xlsx")
_args = _ap.parse_args()
OUT = _args.out
brokers = json.load(open(_args.json, encoding="utf-8"))

NA = "غير متوفرة"
FONT = "Arial"
MAIN, ORIG, LISTS, GUIDE, SUMMARY = "الشركات", "الأصل", "القوائم", "دليل الاستخدام", "ملخص"
LAST = 1000  # formulas and validation reach this row so new companies are covered

# ---- lists -----------------------------------------------------------------
VERIFY = ["موثّقة", "موثّقة جزئياً", "غير موثّقة", "قيد التحقّق", "مرفوضة"]
DATA_ST = ["غير معدّلة", "تم التعديل", "شركة جديدة"]
YESNO = ["نعم", "لا", NA]
PARTNER = ["شريك معتمد", "غير متعاقد"]
PUBLISH = ["منشورة", "مسودة"]

REG = {"fca": "FCA", "cysec": "CySEC", "asic": "ASIC", "fsca": "FSCA", "dfsa": "DFSA",
       "fsa": "FSA", "cbcs": "CBCS", "fscm": "FSC"}
PLAT = {"mt4": "MT4", "mt5": "MT5", "ctrader": "cTrader", "tradingview": "TradingView",
        "proprietary": "منصّة خاصة"}
BADGE = {"hot": "الأعلى طلباً", "best_welcome": "أفضل بونص ترحيبي", "platinum": "شريك بلاتيني",
         "top_rated": "الأعلى تقييماً", "low_spread": "أقل سبريد"}

def v(x):
    """Our value, or 'غير متوفرة' when the site has nothing for it."""
    if x is None:
        return NA
    if isinstance(x, str):
        return x.strip() or NA
    if isinstance(x, list):
        return "، ".join(x) if x else NA
    return x


def yn(b):
    # Tri-state (migration 0033): true / false / null = not checked.
    return "نعم" if b is True else "لا" if b is False else NA


VERIFY_LABEL = {"verified": "موثّقة", "partial": "موثّقة جزئياً", "unverified": "غير موثّقة",
                "in_review": "قيد التحقّق", "rejected": "مرفوضة"}


def verify_status(b):
    if b.get("verification_status") in VERIFY_LABEL:
        return VERIFY_LABEL[b["verification_status"]]
    return "غير موثّقة"


def first_link(b, key):
    links = b.get("broker_links") or []
    return v(links[0].get(key)) if links else NA


# ---- columns: (group, header, width, kind, getter, fmt) --------------------
# kind: key (do not edit) · edit · list:<name> · auto (formula) · ro (read-only) · ext
def ext(key):
    return lambda b: v((b["external_data"] or {}).get(key))


COLS = [
    ("المتابعة", "#", 5, "ro", None, "0"),
    ("المتابعة", "المعرّف (slug)", 18, "key", lambda b: b["slug"], None),
    ("المتابعة", "اسم الشركة", 24, "edit", lambda b: b["name"], None),
    ("المتابعة", "حالة التوثيق", 15, "list:verify", verify_status, None),
    ("المتابعة", "حالة البيانات", 14, "auto", None, None),
    ("المتابعة", "حقول معدّلة", 10, "auto", None, "0"),
    ("المتابعة", "نسبة الاكتمال", 11, "auto", None, "0%"),
    ("المتابعة", "المسؤول", 14, "edit", lambda b: "", None),
    ("المتابعة", "تاريخ آخر مراجعة", 14, "edit", lambda b: date.fromisoformat(b["verified_at"]) if b.get("verified_at") else None, "yyyy-mm-dd"),

    ("الظهور في الموقع", "النشر", 11, "list:publish", lambda b: "منشورة" if b["is_published"] else "مسودة", None),
    ("الظهور في الموقع", "علاقة الشراكة", 13, "list:partner", lambda b: "شريك معتمد" if b["status"] == "partnered" else "غير متعاقد", None),
    ("الظهور في الموقع", "ترتيب العرض", 10, "edit", lambda b: b["sort_order"], "0"),
    ("الظهور في الموقع", "رابط الشعار", 28, "edit", lambda b: v(b["logo_url"]), None),
    ("الظهور في الموقع", "الوصف في الموقع", 55, "edit", lambda b: v(b["description"]), None),

    ("الترخيص والكيان القانوني", "التراخيص الموثّقة", 18, "edit", lambda b: v([REG.get(x, x) for x in b["licenses"]]), None),
    ("الترخيص والكيان القانوني", "الكيان القانوني", 34, "edit", lambda b: v(b.get("legal_entity")), None),
    ("الترخيص والكيان القانوني", "رقم الترخيص", 26, "edit", lambda b: v(b.get("licence_numbers")), None),
    ("الترخيص والكيان القانوني", "رابط التحقّق في سجل الجهة الرقابية", 30, "edit", lambda b: v(b.get("verification_url")), None),
    ("الترخيص والكيان القانوني", "الموقع الرسمي", 26, "edit", lambda b: v(b.get("official_website")), None),

    ("شروط التداول", "سبريد يبدأ من (نقطة)", 12, "edit", lambda b: v(b["spread_from"]), "0.0#"),
    ("شروط التداول", "الرافعة القصوى", 12, "edit", lambda b: v(b["leverage_max"]), None),
    ("شروط التداول", "الحد الأدنى للإيداع ($)", 13, "edit", lambda b: v(b["min_deposit"]), "#,##0"),
    ("شروط التداول", "المنصّات", 18, "edit", lambda b: v([PLAT.get(x, x) for x in b["platforms"]]), None),
    ("شروط التداول", "طرق الإيداع", 20, "edit", lambda b: v(b["deposit_methods"]), None),
    ("شروط التداول", "الدول المقبولة", 18, "edit", lambda b: v(b["accepted_countries"]), None),

    ("المزايا", "حساب إسلامي (بدون سواب)", 12, "list:yesno", lambda b: yn(b["swap_free"]), None),
    ("المزايا", "الروبوتات (EA)", 11, "list:yesno", lambda b: yn(b["supports_ea"]), None),
    ("المزايا", "التحوّط (Hedging)", 11, "list:yesno", lambda b: yn(b["allows_hedging"]), None),
    ("المزايا", "السكالبينغ", 11, "list:yesno", lambda b: yn(b["allows_scalping"]), None),
    ("المزايا", "تداول الذهب", 11, "list:yesno", lambda b: yn(b["supports_gold"]), None),

    ("البونصات", "بونص الإيداع", 13, "edit", lambda b: v(b["deposit_bonus"]), None),
    ("البونصات", "بونص ترحيبي", 13, "edit", lambda b: v(b["welcome_bonus"]), None),
    ("البونصات", "بونص بدون إيداع", 11, "list:yesno", lambda b: yn(b["bonus_no_deposit"]), None),
    ("البونصات", "البونص قابل للسحب", 11, "list:yesno", lambda b: yn(b["bonus_withdrawable"]), None),
    ("البونصات", "الشارات", 16, "edit", lambda b: v([BADGE.get(x, x) for x in b["badges"]]), None),

    ("الشراكة والإحالة", "رابط الإحالة", 28, "edit", lambda b: v([l["referral_url"] for l in b.get("broker_links") or []]), None),
    ("الشراكة والإحالة", "عمولة الوكيل", 15, "edit", lambda b: first_link(b, "agent_commission"), None),
    ("الشراكة والإحالة", "مزايا العميل", 18, "edit", lambda b: first_link(b, "client_benefits"), None),

    ("التقييمات — للقراءة فقط", "تقييم العملاء (من 5)", 11, "ro", lambda b: b["rating"] if b["reviews_count"] > 0 and b["rating"] > 0 else NA, "0.0"),
    ("التقييمات — للقراءة فقط", "عدد المراجعات", 10, "ro", lambda b: b["reviews_count"], "0"),
    ("التقييمات — للقراءة فقط", "التقييم الخارجي (من 10)", 11, "ro", lambda b: v(b["external_score"]), "0.0"),

    ("مصدر خارجي — غير موثّق، للاسترشاد فقط", "الجهات الرقابية", 18, "ext", ext("regulation"), None),
    ("مصدر خارجي — غير موثّق، للاسترشاد فقط", "تفاصيل التراخيص", 30, "ext", ext("regulation_details"), None),
    ("مصدر خارجي — غير موثّق، للاسترشاد فقط", "بلد التنظيم", 14, "ext", ext("regulation_country"), None),
    ("مصدر خارجي — غير موثّق، للاسترشاد فقط", "الرافعة المالية", 12, "ext", ext("leverage"), None),
    ("مصدر خارجي — غير موثّق، للاسترشاد فقط", "الحد الأدنى للإيداع", 12, "ext", ext("min_deposit"), None),
    ("مصدر خارجي — غير موثّق، للاسترشاد فقط", "سنة التأسيس", 10, "ext", ext("founded_year"), None),
    ("مصدر خارجي — غير موثّق، للاسترشاد فقط", "سنوات العمل", 11, "ext", ext("years"), None),
    ("مصدر خارجي — غير موثّق، للاسترشاد فقط", "المقرّ", 16, "ext", ext("headquarters"), None),
    ("مصدر خارجي — غير موثّق، للاسترشاد فقط", "نموذج التنفيذ", 14, "ext", ext("business_model"), None),
    ("مصدر خارجي — غير موثّق، للاسترشاد فقط", "نوع الحساب", 14, "ext", ext("account_type"), None),
    ("مصدر خارجي — غير موثّق، للاسترشاد فقط", "المنصّات", 18, "ext", ext("platforms"), None),
    ("مصدر خارجي — غير موثّق، للاسترشاد فقط", "الموقع الرسمي", 26, "ext", ext("website_url"), None),

    ("ملاحظات", "ملاحظات", 36, "edit", lambda b: "", None),
]

NOTES = {
    "المعرّف (slug)": "مفتاح الربط مع الموقع — لا تعدّله. يُستخدم لإعادة رفع البيانات.",
    "حالة التوثيق": "اختر من القائمة. تُعدَّل يدوياً حسب نتيجة التحقّق.",
    "حالة البيانات": "تُحسب تلقائياً: تتحوّل إلى «تم التعديل» عند تغيير أي حقل عن النسخة الأصلية.",
    "حقول معدّلة": "عدد الحقول التي تغيّرت عن النسخة الأصلية (تلقائي).",
    "نسبة الاكتمال": "نسبة الحقول المعبّأة من أعمدة «الظهور في الموقع» حتى «الشراكة والإحالة» (تلقائي).",
    "التراخيص الموثّقة": "رموز الجهات فقط، مفصولة بفاصلة: FCA، CySEC، ASIC، FSCA، DFSA، FSA، CBCS، FSC",
    "رابط التحقّق في سجل الجهة الرقابية": "رابط صفحة الشركة في سجل الجهة الرقابية نفسها — هذا دليل التوثيق.",
    "تقييم العملاء (من 5)": "يُحسب من مراجعات العملاء الحقيقية فقط — لا يُعدَّل يدوياً.",
    "الجهات الرقابية": "بيانات من دليل خارجي لم تتحقّق منها المنصّة. لا تنقلها إلى الأعمدة الموثّقة قبل التأكد من سجل الجهة الرقابية.",
    "الحد الأدنى للإيداع ($)": "رقم فقط بالدولار، مثل 100",
    "سبريد يبدأ من (نقطة)": "رقم فقط بالنقاط، مثل 0.6",
}

# ---- styles ----------------------------------------------------------------
GROUP_FILL = {
    "المتابعة": "1F2A44", "الظهور في الموقع": "2E4A7D", "الترخيص والكيان القانوني": "1E6B52",
    "شروط التداول": "2F6F8F", "المزايا": "3F6E8C", "البونصات": "7A5C1E", "الشراكة والإحالة": "5B3F7A",
    "التقييمات — للقراءة فقط": "5F6B7A", "مصدر خارجي — غير موثّق، للاسترشاد فقط": "8A4B2A", "ملاحظات": "444444",
}
HEAD_FILL = {"key": "D9DDE3", "ro": "E4E7EB", "ext": "F6E3D6", "auto": "DCE6F2"}
thin = Side(style="thin", color="D0D5DD")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=False, readingOrder=2)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2)
f = lambda **k: Font(name=FONT, size=k.pop("size", 10), **k)


def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


wb = Workbook()
ws = wb.active
ws.title = MAIN
orig = wb.create_sheet(ORIG)
for sh in (ws, orig):
    sh.sheet_view.rightToLeft = True

ncol = len(COLS)
# explicit positions we need in formulas
C_SLUG, C_STATUS, C_CHANGED, C_COMPLETE = 2, 5, 6, 7
FIRST_DATA = 10                       # «النشر»
LAST_FILL = next(i for i, c in enumerate(COLS, 1) if c[1] == "مزايا العميل")
LAST_CMP = ncol - 1                   # everything except notes is compared to the original
L_FD, L_LF, L_LC = L(FIRST_DATA), L(LAST_FILL), L(LAST_CMP)

# group band (row 1) and headers (row 2)
start = 1
for i in range(1, ncol + 2):
    g = COLS[i - 1][0] if i <= ncol else None
    if i > ncol or g != COLS[start - 1][0]:
        for sh in (ws, orig):
            if i - 1 > start:
                sh.merge_cells(start_row=1, start_column=start, end_row=1, end_column=i - 1)
            c = sh.cell(1, start, COLS[start - 1][0])
            c.font = f(bold=True, color="FFFFFF", size=11)
            c.fill = fill(GROUP_FILL[COLS[start - 1][0]])
            c.alignment = CENTER
        start = i

for i, (g, name, width, kind, _, fmt) in enumerate(COLS, 1):
    for sh in (ws, orig):
        c = sh.cell(2, i, name)
        c.font = f(bold=True, color="1F2A44")
        c.fill = fill(HEAD_FILL.get(kind, "EEF2F7"))
        c.alignment = CENTER
        c.border = BORDER
        sh.column_dimensions[L(i)].width = width
    if name in NOTES and not (kind == "ext" and name != "الجهات الرقابية"):
        ws.cell(2, i).comment = Comment(NOTES[name], "FX Partners", width=260, height=90)

ws.row_dimensions[1].height = 24
ws.row_dimensions[2].height = 44

# data
for r, b in enumerate(brokers, 3):
    for i, (g, name, width, kind, get, fmt) in enumerate(COLS, 1):
        if i == 1:
            val = f"=ROW()-2"
        elif i == C_STATUS:
            val = (f'=IF(${L(C_SLUG)}{r}="","",IFERROR(IF(MATCH(${L(C_SLUG)}{r},\'{ORIG}\'!$B$1:$B${LAST},0)>0,'
                   f'IF({L(C_CHANGED)}{r}>0,"تم التعديل","غير معدّلة")),"شركة جديدة"))')
        elif i == C_CHANGED:
            val = (f'=IF(${L(C_SLUG)}{r}="","",IFERROR(SUMPRODUCT(--({L_FD}{r}:{L_LC}{r}<>'
                   f'INDEX(\'{ORIG}\'!${L_FD}$1:${L_LC}${LAST},MATCH(${L(C_SLUG)}{r},\'{ORIG}\'!$B$1:$B${LAST},0),0))),0))')
        elif i == C_COMPLETE:
            val = (f'=IF(${L(C_SLUG)}{r}="","",(COUNTA({L_FD}{r}:{L_LF}{r})-COUNTIF({L_FD}{r}:{L_LF}{r},"{NA}"))'
                   f'/COLUMNS({L_FD}{r}:{L_LF}{r}))')
        else:
            val = get(b)
        c = ws.cell(r, i, val)
        c.font = f(color="7A4A2A" if kind == "ext" else ("5F6B7A" if kind in ("key", "ro", "auto") else "111111"))
        c.alignment = Alignment(horizontal="center" if kind.startswith(("list", "auto", "ro")) or i == 1 else "right",
                                vertical="center", readingOrder=2)
        c.border = BORDER
        if fmt:
            c.number_format = fmt
        if kind == "key":
            c.fill = fill("F2F4F7")
        # snapshot of the original (compared against to detect edits)
        # Name, verification status and review date are copied too, so the
        # importer can tell whether someone changed them.
        if i >= FIRST_DATA or i in (C_SLUG, 3, 4, 9):
            orig.cell(r, i, val)
    ws.row_dimensions[r].height = 22

last_row = len(brokers) + 2
ws.freeze_panes = "D3"
ws.auto_filter.ref = f"A2:{L(ncol)}{last_row}"
orig.sheet_state = "hidden"

# pre-format rows for companies added later
for r in range(last_row + 1, last_row + 51):
    for i, c in enumerate(COLS, 1):
        cell = ws.cell(r, i)
        cell.border = BORDER
        cell.font = f()
        if c[5]:
            cell.number_format = c[5]
        if i == 1:
            cell.value = f'=IF($B{r}="","",ROW()-2)'
        elif i == C_STATUS:
            cell.value = f'=IF($B{r}="","","شركة جديدة")'
        elif i == C_COMPLETE:
            cell.value = (f'=IF($B{r}="","",(COUNTA({L_FD}{r}:{L_LF}{r})-COUNTIF({L_FD}{r}:{L_LF}{r},"{NA}"))'
                          f'/COLUMNS({L_FD}{r}:{L_LF}{r}))')

# ---- lists sheet + validation ----------------------------------------------
ls = wb.create_sheet(LISTS)
ls.sheet_view.rightToLeft = True
lists = {"verify": ("حالة التوثيق", VERIFY), "yesno": ("نعم / لا", YESNO),
         "partner": ("علاقة الشراكة", PARTNER), "publish": ("النشر", PUBLISH), "data": ("حالة البيانات", DATA_ST)}
ref = {}
for j, (key, (title, items)) in enumerate(lists.items(), 1):
    h = ls.cell(1, j, title)
    h.font = f(bold=True, color="FFFFFF")
    h.fill = fill("1F2A44")
    h.alignment = CENTER
    for k, it in enumerate(items, 2):
        ls.cell(k, j, it).font = f()
    ls.column_dimensions[L(j)].width = 18
    ref[key] = f"'{LISTS}'!${L(j)}$2:${L(j)}${len(items) + 1}"
ls.cell(len(max(lists.values(), key=lambda x: len(x[1]))[1]) + 3, 1,
        "يمكنك إضافة خيارات جديدة أسفل أي قائمة ثم توسيع مداها من «التحقّق من صحة البيانات».").font = f(italic=True, color="5F6B7A")

for i, c in enumerate(COLS, 1):
    if c[3].startswith("list:"):
        key = c[3].split(":")[1]
        dv = DataValidation(type="list", formula1="=" + ref[key], allow_blank=True, showDropDown=False)
        dv.error, dv.errorTitle = "اختر قيمة من القائمة", "قيمة غير صالحة"
        dv.prompt, dv.promptTitle = "اختر من القائمة", c[1]
        ws.add_data_validation(dv)
        dv.add(f"{L(i)}3:{L(i)}{LAST}")
dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(2020,1,1)", allow_blank=True)
dv.error = "أدخل تاريخاً بصيغة 2026-09-24"
ws.add_data_validation(dv)
dv.add(f"I3:I{LAST}")

# ---- conditional formatting ------------------------------------------------
full = f"A3:{L(ncol)}{LAST}"
# edited cells (compared with the hidden original, by slug, so sorting is safe)
ws.conditional_formatting.add(
    f"{L_FD}3:{L_LC}{LAST}",
    FormulaRule(formula=[f'AND($B3<>"",IFERROR({L_FD}3<>INDEX(INDIRECT("\'{ORIG}\'!"&ADDRESS(1,COLUMN(),4)&":"&ADDRESS({LAST},COLUMN(),4)),'
                         f'MATCH($B3,INDIRECT("\'{ORIG}\'!B1:B{LAST}"),0)),FALSE))'],
                fill=fill("FFF3B0"), font=Font(name=FONT, bold=True, color="7A5C00"), stopIfTrue=True))
ws.conditional_formatting.add(f"{L_FD}3:{L(ncol)}{LAST}",
                              CellIsRule(operator="equal", formula=[f'"{NA}"'], fill=fill("F4F5F7"), font=Font(name=FONT, color="A0A7B2", italic=True)))
colors = {"موثّقة": ("D1FADF", "05603A"), "موثّقة جزئياً": ("E0F2E9", "1E6B52"), "غير موثّقة": ("FEF0C7", "93370D"),
          "قيد التحقّق": ("D1E9FF", "1849A9"), "مرفوضة": ("FEE4E2", "B42318"),
          "تم التعديل": ("D1FADF", "05603A"), "غير معدّلة": ("F2F4F7", "475467"), "شركة جديدة": ("D1E9FF", "1849A9")}
for col in (L(4), L(C_STATUS)):
    for text, (bg, fg) in colors.items():
        ws.conditional_formatting.add(f"{col}3:{col}{LAST}",
                                      CellIsRule(operator="equal", formula=[f'"{text}"'], fill=fill(bg), font=Font(name=FONT, bold=True, color=fg)))
for text, (bg, fg) in {"نعم": ("D1FADF", "05603A"), "لا": ("FEE4E2", "B42318"), "مسودة": ("FEF0C7", "93370D")}.items():
    ws.conditional_formatting.add(f"{L_FD}3:{L_LF}{LAST}",
                                  CellIsRule(operator="equal", formula=[f'"{text}"'], fill=fill(bg), font=Font(name=FONT, color=fg)))
ws.conditional_formatting.add(f"{L(C_COMPLETE)}3:{L(C_COMPLETE)}{LAST}",
                              ColorScaleRule(start_type="num", start_value=0, start_color="FDE2E1",
                                             mid_type="num", mid_value=0.5, mid_color="FEF3C7",
                                             end_type="num", end_value=1, end_color="C6F6D5"))

# ---- summary sheet ---------------------------------------------------------
sm = wb.create_sheet(SUMMARY, 0)
sm.sheet_view.rightToLeft = True
sm.column_dimensions["A"].width = 34
sm.column_dimensions["B"].width = 14
sm.column_dimensions["C"].width = 14
sm.column_dimensions["D"].width = 4
sm.column_dimensions["E"].width = 34
sm.column_dimensions["F"].width = 14
sm.column_dimensions["G"].width = 14
t = sm.cell(1, 1, "FX Partners — ملخص بيانات الشركات")
t.font = f(bold=True, size=16, color="1F2A44")
sm.merge_cells("A1:G1")
sm.merge_cells("A2:G2")
sm.row_dimensions[1].height = 28
sm.cell(2, 1, f"أُنشئ من قاعدة بيانات الموقع بتاريخ {date.today():%Y-%m-%d} — الأرقام أدناه تتحدّث تلقائياً مع كل تعديل.").font = f(color="5F6B7A")
M = f"'{MAIN}'!"
rng = lambda col: f"{M}${col}$3:${col}${LAST}"


def block(row, col, title, items, key_col):
    h1 = sm.cell(row, col, title)
    h2 = sm.cell(row, col + 1, "العدد")
    h3 = sm.cell(row, col + 2, "النسبة")
    for h in (h1, h2, h3):
        h.font = f(bold=True, color="FFFFFF")
        h.fill = fill("1F2A44")
        h.alignment = CENTER
    for k, it in enumerate(items, row + 1):
        a = sm.cell(k, col, it)
        a.font = f()
        a.alignment = RIGHT
        n = sm.cell(k, col + 1, f'=COUNTIF({rng(key_col)},{L(col)}{k})')
        p = sm.cell(k, col + 2, f'=IF($B$5=0,0,{L(col + 1)}{k}/$B$5)')
        p.number_format = "0%"
        for c in (a, n, p):
            c.border = BORDER
        n.alignment = p.alignment = Alignment(horizontal="center")
    return row + len(items) + 2


r = 4
for k, (label, formula, fmt) in enumerate([
    ("إجمالي الشركات", f'=COUNTA({rng("B")})', "0"),
    ("منشورة في الموقع", f'=COUNTIF({rng(L(FIRST_DATA))},"منشورة")', "0"),
    ("متوسط نسبة الاكتمال", f'=IFERROR(AVERAGE({rng(L(C_COMPLETE))}),0)', "0%"),
    ("شركات عُدّلت بياناتها", f'=COUNTIF({rng(L(C_STATUS))},"تم التعديل")', "0"),
], 5):
    a = sm.cell(k, 1, label)
    b = sm.cell(k, 2, formula)
    a.font = f(bold=True)
    b.font = f(bold=True, size=12, color="1F2A44")
    b.number_format = fmt
    b.alignment = Alignment(horizontal="center")
    for c in (a, b):
        c.border = BORDER
        c.fill = fill("EEF2F7")
sm.cell(4, 1, "نظرة عامة").font = f(bold=True, size=12, color="1F2A44")

r = block(10, 1, "حالة التوثيق", VERIFY, "D")
r = block(r, 1, "حالة البيانات", DATA_ST, L(C_STATUS))

# field completeness
h = [sm.cell(10, 5, "الحقل"), sm.cell(10, 6, "معبّأ"), sm.cell(10, 7, "نسبة الاكتمال")]
for c in h:
    c.font = f(bold=True, color="FFFFFF")
    c.fill = fill("1F2A44")
    c.alignment = CENTER
rr = 11
for i in range(FIRST_DATA, LAST_FILL + 1):
    if COLS[i - 1][1] in ("النشر", "علاقة الشراكة", "ترتيب العرض"):
        continue
    col = L(i)
    a = sm.cell(rr, 5, COLS[i - 1][1])
    n = sm.cell(rr, 6, f'=COUNTIFS({rng("B")},"<>",{rng(col)},"<>{NA}",{rng(col)},"<>")')
    p = sm.cell(rr, 7, f'=IF($B$5=0,0,F{rr}/$B$5)')
    p.number_format = "0%"
    a.font = n.font = p.font = f()
    n.alignment = p.alignment = Alignment(horizontal="center")
    for c in (a, n, p):
        c.border = BORDER
    rr += 1
sm.conditional_formatting.add(f"G11:G{rr - 1}", ColorScaleRule(start_type="num", start_value=0, start_color="FDE2E1",
                                                               mid_type="num", mid_value=0.5, mid_color="FEF3C7",
                                                               end_type="num", end_value=1, end_color="C6F6D5"))
sm.conditional_formatting.add("C11:C40", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                                        end_type="num", end_value=1, end_color="C7D7FE"))

# ---- guide sheet -----------------------------------------------------------
gd = wb.create_sheet(GUIDE, 0)
gd.sheet_view.rightToLeft = True
gd.sheet_view.showGridLines = False
gd.column_dimensions["A"].width = 3
gd.column_dimensions["B"].width = 30
gd.column_dimensions["C"].width = 62
gd.column_dimensions["D"].width = 26
row = 2
gd.cell(row, 2, "دليل استخدام ملف بيانات الشركات").font = f(bold=True, size=18, color="1F2A44")
row += 1
gd.cell(row, 2, "الهدف: جمع المعلومات الحقيقية لكل شركة يدوياً من مصادرها الرسمية، ثم إعادة رفعها إلى الموقع.").font = f(color="5F6B7A", size=11)
row += 2


def section(title):
    global row
    c = gd.cell(row, 2, title)
    c.font = f(bold=True, size=13, color="FFFFFF")
    c.fill = fill("1F2A44")
    gd.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    gd.row_dimensions[row].height = 22
    row += 1


def line(a, b="", c="", bg=None, fg="111111", bold=False):
    global row
    for j, val in enumerate((a, b, c), 2):
        cell = gd.cell(row, j, val)
        cell.font = f(color=fg, bold=bold or j == 2)
        cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True, readingOrder=2)
        cell.border = BORDER
    if not c:
        gd.merge_cells(start_row=row, start_column=3, end_row=row, end_column=4)
    if bg:
        gd.cell(row, 2).fill = fill(bg)
    gd.row_dimensions[row].height = 32
    row += 1


section("خطوات العمل")
for i, s in enumerate([
    "افتح ورقة «الشركات». كل صف شركة، وكل عمود معلومة يعرضها الموقع.",
    "ابحث عن الشركة في موقعها الرسمي وفي سجل الجهة الرقابية (FCA، CySEC، ASIC…).",
    "اكتب المعلومة الحقيقية مكان «غير متوفرة». اترك «غير متوفرة» لكل ما لم تتأكّد منه.",
    "ضع رابط صفحة الشركة في سجل الجهة الرقابية في عمود «رابط التحقّق» — هذا دليل التوثيق.",
    "غيّر «حالة التوثيق» من القائمة، واكتب اسمك وتاريخ المراجعة.",
    "«حالة البيانات» تتغيّر وحدها إلى «تم التعديل»، وتظهر الخلايا المعدّلة بخلفية صفراء.",
    "أرسل الملف بعد الانتهاء: يتحوّل إلى ملف SQL يُكتب فيه ما تغيّر فقط (المفتاح هو عمود المعرّف slug).",
], 1):
    line(f"الخطوة {i}", s)
row += 1

section("دلالة الألوان")
line("غير متوفرة", "المعلومة غير موجودة في الموقع حالياً — المطلوب تعبئتها.", "", "F4F5F7", "A0A7B2")
line("خلية معدّلة", "قيمة تغيّرت عن النسخة الأصلية المأخوذة من الموقع.", "", "FFF3B0", "7A5C00")
line("أعمدة بنية", "بيانات من دليل خارجي غير موثّق — للاسترشاد فقط، لا تُعتمد قبل التحقّق.", "", "F6E3D6", "7A4A2A")
line("أعمدة رمادية", "تلقائية أو للقراءة فقط (المعرّف، الحالة، الاكتمال، التقييمات).", "", "E4E7EB", "5F6B7A")
row += 1

section("حالات التوثيق")
for s, d in [("موثّقة", "تم التحقّق من الكيان القانوني ورقم الترخيص في سجل الجهة الرقابية."),
             ("موثّقة جزئياً", "بعض البيانات متحقَّق منها (مثل الترخيص) وبقية الشروط لم تُراجَع."),
             ("غير موثّقة", "البيانات من مصدر خارجي ولم يتحقّق منها أحد بعد."),
             ("قيد التحقّق", "المراجعة جارية الآن."),
             ("مرفوضة", "تعذّر التحقّق أو الشركة غير مرخّصة — لا تُنشر.")]:
    bg, fg = colors[s]
    line(s, d, "", bg, fg)
row += 1

section("صيغة الإدخال لكل حقل")
line("الحقل", "ما المطلوب", "صيغة الإدخال", "EEF2F7", "1F2A44", True)
for a, b, c in [
    ("التراخيص الموثّقة", "الجهات الرقابية التي تغطّي حساب العميل فعلاً", "FCA، CySEC"),
    ("الكيان القانوني", "اسم الشركة المسجّلة لدى الجهة الرقابية", "اسم الكيان (الجهة)"),
    ("رقم الترخيص", "رقم الترخيص لكل جهة", "FCA: 123456"),
    ("سبريد يبدأ من", "أقل سبريد معلن بالنقاط", "رقم: 0.6"),
    ("الرافعة القصوى", "أعلى رافعة متاحة", "1:500"),
    ("الحد الأدنى للإيداع", "بالدولار الأمريكي", "رقم: 100"),
    ("المنصّات", "المنصّات المتاحة", "MT4، MT5، cTrader"),
    ("طرق الإيداع", "وسائل الإيداع", "تحويل بنكي، بطاقات، عملات رقمية"),
    ("الدول المقبولة", "رموز الدول بحرفين", "IQ، SA، AE"),
    ("المزايا والبونص (نعم/لا)", "«لا» تعني: تحقّقت ولم تجدها. إن لم تتحقّق اترك «غير متوفرة»", "نعم / لا / غير متوفرة"),
    ("تاريخ آخر مراجعة", "تاريخ التحقّق", "2026-09-24"),
]:
    line(a, b, c)
row += 1

section("قواعد مهمّة")
for s in [
    "لا تعدّل عمود «المعرّف (slug)» — هو ما يربط الصف بصفحة الشركة في الموقع.",
    "تقييم العملاء وعدد المراجعات يُحسبان من مراجعات حقيقية فقط، ولا يُكتبان يدوياً.",
    "لا تنقل بيانات الأعمدة البنية (المصدر الخارجي) إلى الأعمدة الموثّقة قبل التأكّد منها في المصدر الرسمي.",
    "اكتب الترخيص الذي يغطّي حساب العميل الذي يُفتح عبر رابط إحالتك، لا كل تراخيص المجموعة.",
    "لإضافة شركة جديدة: اكتبها في أول صف فارغ أسفل الجدول مع معرّف (slug) جديد بالإنجليزية.",
    "الورقة المخفية «الأصل» تحفظ النسخة الأصلية للمقارنة — لا تحذفها.",
]:
    line("•", s)

for sh, color in ((gd, "1F2A44"), (sm, "2E4A7D"), (ws, "1E6B52"), (ls, "8A94A6")):
    sh.sheet_properties.tabColor = color
for sh in (ws, sm, gd):
    sh.page_setup.orientation = "landscape"
    sh.page_setup.fitToWidth = 1
    sh.page_setup.fitToHeight = 0
    sh.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_title_rows = "1:2"

wb.active = wb.sheetnames.index(MAIN)
wb.save(OUT)
print(OUT, len(brokers), "rows,", ncol, "columns; first data col", L_FD, "last fill", L_LF, "last cmp", L_LC)
